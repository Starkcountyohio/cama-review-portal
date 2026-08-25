"""
monthly_retrain.py — Stark County Auditor
Monthly AI Grade & Condition Prompt Retraining

Run when Jason says: python monthly_retrain.py
  Add --prompt-only to retrain the prompt in the template WITHOUT rebuilding or
  pushing the portal (the new prompt then ships with next week's normal build).

What it does:
  1. Loads the 4 most recent weekly archive JSON files (≈ one month)
  2. Queries Oracle DWELDAT for current Grade/CDU on those parcels
  3. Saves Grade&Con_Compara.xlsx (current month snapshot)
  4. Compares before (portal data at time of review) vs after (appraiser decisions)
  5. Prints full findings report — what changed, what signals moved
  6. Rewrites the AI prompt in review_portal_template.html
  7. Rebuilds review_portal.html with current week's data
  8. Pushes to GitHub
"""

import sys
import json
import math
import subprocess
import statistics
from collections import Counter
from datetime import date, timedelta
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

try:
    import oracledb
    import pandas as pd
    from openpyxl import Workbook
except ImportError as e:
    print(f"Missing dependency: {e}  — run: pip install oracledb pandas openpyxl")
    sys.exit(1)

# ── Paths ──────────────────────────────────────────────────────────────────────
ROOT       = Path(__file__).parent
ARCHIVE    = ROOT / "archive"
TEMPLATE   = ROOT / "review_portal_template.html"
COMPARA    = ROOT / "Grade&Con_Compara.xlsx"
BUILD      = ROOT / "build_portal.py"
AD_DIR     = ROOT / "Automation Domination"

sys.path.insert(0, str(AD_DIR))
from credentials import ORACLE_USER, ORACLE_PASSWORD, ORACLE_DSN
# Single source of truth for the future-year layer (same import make_photo_upload uses;
# playwright is imported lazily inside create_future_year, so this is cheap).
from create_future_year import TARGET_TAX_YEAR

# ── Grade/CDU rankings ─────────────────────────────────────────────────────────
CDU_RANK = {"PR": 0, "FR": 1, "AV": 2, "GD": 3, "VG": 4, "EX": 5}

# ── Keyword groups ─────────────────────────────────────────────────────────────
UPGRADE_PHRASES = [
    "new roof", "newer roof", "new furnace", "new hvac", "new a/c", "new windows",
    "updated kitchen", "updated bath", "renovated", "remodeled", "new floors",
    "new electric", "new plumbing", "new water heater", "new siding",
    "move-in ready", "move in ready", "updated throughout", "completely updated",
    "fully updated", "immaculate", "meticulously", "turnkey",
]
DOWNGRADE_PHRASES = [
    "tlc", "opportunity", "as-is", "as is", "investor", "handyman", "needs work",
    "needs updating", "fixer", "sweat equity", "elbow grease", "potential",
    "estate sale", "cash only", "sold as is", "sold as-is", "vacant", "abandoned",
]
NOISE_PHRASES = [
    "great location", "convenient location", "minutes from", "close to",
    "near shopping", "desirable", "sought after", "award winning school",
    "location location",
]
HARD_NEG = ["estate sale", "cash only", "sold as is", "sold as-is", "vacant", "abandoned"]


# ── Step 1: Load recent archive JSON files ─────────────────────────────────────

def load_recent_archives(n_weeks=4, skip_recent=0):
    """Load the n most recent weekly archive JSON files.

    skip_recent drops that many of the NEWEST weeks first. The current week's
    parcels have no appraiser decisions yet — Jason builds the portal Monday and
    appraisers review Tue–Fri — so including the freshest week adds a block of
    guaranteed no-change rows that reads as "appraiser agreed" and dilutes every
    rate. Use skip_recent=1 when retraining on the same day as a build.
    """
    files = sorted(ARCHIVE.glob("*.json"))
    if not files:
        print("ERROR: No archive JSON files found in archive/")
        sys.exit(1)
    if skip_recent:
        skipped = [f.stem for f in files[-skip_recent:]]
        files = files[:-skip_recent]
        print(f"  Skipping newest {skip_recent} week(s) (not yet reviewed): "
              f"{', '.join(skipped)}")
    recent = files[-n_weeks:]
    print(f"  Archive files: {', '.join(f.stem for f in recent)}")

    parcels = {}
    for jf in recent:
        with open(jf, encoding="utf-8") as f:
            data = json.load(f)
        rows = data.get("mismatches", []) + data.get("perfects", [])
        for p in rows:
            pid = str(p.get("Parcel_ID", "")).strip()
            if pid and pid not in parcels:
                parcels[pid] = dict(p, _week=jf.stem)

    print(f"  Unique parcels loaded: {len(parcels)}")
    return parcels, [f.stem for f in recent]


# ── Step 2: Query Oracle for current Grade/CDU ─────────────────────────────────

def fetch_grade_cdu(parcel_ids: list) -> dict:
    """Query DWELDAT for the appraiser's CURRENT GRADE/CDU. Returns {parid: row}.

    Reads BOTH the current-year and future-year layers and prefers the future year
    when a row exists there. Since create_future_year.py went to prod (2026-W29),
    appraisers record their Grade/CDU decisions on the FUTURE-year record, not the
    current one — sales are being ordered into the future year while CAMA is still
    on the current one. Querying only `TAXYR = current_year` therefore reads a layer
    nobody has touched and reports ~0% change for every recent week, which silently
    deflates every rate and signal weight in the report. Measured 2026-08-25 over
    W24–W35: 0–1% change vs the current-year layer for W29+, 34–41% vs the future
    year, with the break falling exactly on the week the 2027 workflow started.
    Older weeks have no future-year row and fall back to the current year.
    """
    current_year = date.today().year
    print(f"  Connecting to Oracle ({ORACLE_DSN})...")
    print(f"  Layers: preferring TAXYR={TARGET_TAX_YEAR}, falling back to {current_year}")

    conn = oracledb.connect(user=ORACLE_USER, password=ORACLE_PASSWORD, dsn=ORACLE_DSN)
    by_year = {}   # parid -> {taxyr: row}

    # Oracle IN clause limit is 1000 — batch if needed
    batch_size = 999
    batches = [parcel_ids[i:i+batch_size] for i in range(0, len(parcel_ids), batch_size)]

    try:
        with conn.cursor() as cur:
            for batch in batches:
                placeholders = ",".join(f":{i+1}" for i in range(len(batch)))
                sql = f"""
                    SELECT PARID, GRADE, CDU, TAXYR, CARD
                    FROM DWELDAT
                    WHERE PARID IN ({placeholders})
                      AND TAXYR IN ({current_year}, {TARGET_TAX_YEAR})
                """
                cur.execute(sql, batch)
                for row in cur.fetchall():
                    parid = str(row[0]).strip()
                    by_year.setdefault(parid, {})[int(row[3])] = {
                        "PARID":  parid,
                        "GRADE":  str(row[1]).strip() if row[1] else "",
                        "CDU":    str(row[2]).strip() if row[2] else "",
                        "TAXYR":  row[3],
                        "CARD":   row[4],
                    }
    finally:
        conn.close()

    results, from_future = {}, 0
    for parid, years in by_year.items():
        if TARGET_TAX_YEAR in years:
            results[parid] = years[TARGET_TAX_YEAR]
            from_future += 1
        elif current_year in years:
            results[parid] = years[current_year]

    print(f"  Oracle returned: {len(results)} parcels with Grade/CDU "
          f"({from_future} from {TARGET_TAX_YEAR}, "
          f"{len(results) - from_future} from {current_year})")
    return results


# ── Step 3: Save Grade&Con_Compara.xlsx ───────────────────────────────────────

def save_compara(oracle_data: dict):
    wb = Workbook()
    ws = wb.active
    ws.append(["PARID", "GRADE", "CDU", "TAXYR", "CARD"])
    for row in sorted(oracle_data.values(), key=lambda r: r["PARID"]):
        ws.append([row["PARID"], row["GRADE"], row["CDU"], row["TAXYR"], row["CARD"]])
    wb.save(COMPARA)
    print(f"  Saved: Grade&Con_Compara.xlsx  ({len(oracle_data)} rows)")


# ── Step 4: Build matched dataset ─────────────────────────────────────────────

def build_matched(portal: dict, oracle: dict) -> list:
    matched = []
    unmatched = 0
    for pid, p in portal.items():
        o = oracle.get(pid)
        if not o:
            unmatched += 1
            continue
        orig_grade = str(p.get("GRADE", "")).strip()
        orig_cdu   = str(p.get("CDU", "")).strip()
        fin_grade  = o["GRADE"]
        fin_cdu    = o["CDU"]
        matched.append({
            "parcelId":     pid,
            "orig_grade":   orig_grade,
            "orig_cdu":     orig_cdu,
            "final_grade":  fin_grade,
            "final_cdu":    fin_cdu,
            "grade_changed": orig_grade != fin_grade and fin_grade,
            "cdu_changed":   orig_cdu   != fin_cdu   and fin_cdu,
            "remarks":      str(p.get("Public_Remarks", "")).lower(),
            "yrblt":        p.get("YRBLT"),
            "_week":        p.get("_week", ""),
        })
    if unmatched:
        print(f"  Note: {unmatched} parcels not found in Oracle (may have been removed)")
    return matched


# ── Step 5: Analyze ────────────────────────────────────────────────────────────

def analyze(matched: list) -> dict:
    total      = len(matched)
    grade_chg  = [m for m in matched if m["grade_changed"]]
    cdu_chg    = [m for m in matched if m["cdu_changed"]]
    either_chg = [m for m in matched if m["grade_changed"] or m["cdu_changed"]]
    upgrades   = [m for m in cdu_chg if CDU_RANK.get(m["final_cdu"], 2) > CDU_RANK.get(m["orig_cdu"], 2)]
    downgrades = [m for m in cdu_chg if CDU_RANK.get(m["final_cdu"], 2) < CDU_RANK.get(m["orig_cdu"], 2)]

    cdu_transitions   = Counter(f"{m['orig_cdu']}→{m['final_cdu']}" for m in cdu_chg).most_common(8)
    grade_transitions = Counter(f"{m['orig_grade']}→{m['final_grade']}" for m in grade_chg).most_common(8)

    def phrase_stats(phrases, subset, label):
        stats = []
        for phrase in phrases:
            in_subset = sum(1 for m in subset  if phrase in m["remarks"])
            in_total  = sum(1 for m in matched if phrase in m["remarks"])
            if in_total < 2:
                continue
            baseline   = in_total / total
            rate       = in_subset / len(subset) if subset else 0
            multiplier = rate / baseline if baseline > 0 else 0
            stats.append({
                "phrase": phrase, "count": in_subset, "total_seen": in_total,
                "rate": rate, "baseline": baseline, "multiplier": multiplier,
            })
        return sorted(stats, key=lambda x: -x["multiplier"])

    upgrade_stats   = phrase_stats(UPGRADE_PHRASES,   upgrades,   "upgrade")
    downgrade_stats = phrase_stats(DOWNGRADE_PHRASES, downgrades, "downgrade")

    hard_neg_parcels    = [m for m in matched if any(kw in m["remarks"] for kw in HARD_NEG)]
    hard_neg_change_pct = (sum(1 for m in hard_neg_parcels if m["cdu_changed"])
                           / len(hard_neg_parcels) * 100) if hard_neg_parcels else 0

    noise_in_chg   = sum(1 for m in either_chg if any(kw in m["remarks"] for kw in NOISE_PHRASES))
    noise_in_unch  = sum(1 for m in matched
                         if not (m["grade_changed"] or m["cdu_changed"])
                         and any(kw in m["remarks"] for kw in NOISE_PHRASES))
    n_unchanged = total - len(either_chg)
    noise_chg_pct  = noise_in_chg  / len(either_chg) * 100 if either_chg  else 0
    noise_unch_pct = noise_in_unch / n_unchanged      * 100 if n_unchanged else 0

    yrblts = [m["yrblt"] for m in upgrades
              if isinstance(m.get("yrblt"), (int, float)) and m["yrblt"] and m["yrblt"] > 1800]
    yrblt_median = int(statistics.median(yrblts)) if yrblts else 0

    return {
        "total": total,
        "grade_change_pct":  len(grade_chg)  / total * 100,
        "cdu_change_pct":    len(cdu_chg)    / total * 100,
        "either_change_pct": len(either_chg) / total * 100,
        "n_upgrades":        len(upgrades),
        "n_downgrades":      len(downgrades),
        "cdu_transitions":   cdu_transitions,
        "grade_transitions": grade_transitions,
        "upgrade_stats":     upgrade_stats[:12],
        "downgrade_stats":   downgrade_stats[:10],
        "hard_neg_n":           len(hard_neg_parcels),
        "hard_neg_change_pct":  hard_neg_change_pct,
        "noise_chg_pct":        noise_chg_pct,
        "noise_unch_pct":       noise_unch_pct,
        "yrblt_median_upgraded": yrblt_median,
    }


# ── Step 6: Print findings report ─────────────────────────────────────────────

def print_report(f: dict, weeks: list):
    print("\n" + "=" * 65)
    print("  MONTHLY RETRAINING REPORT")
    print("=" * 65)
    print(f"  Weeks analyzed:      {', '.join(weeks)}")
    print(f"  Parcels matched:     {f['total']}")
    print()
    print("  CHANGE RATES:")
    print(f"    CDU changed:       {f['cdu_change_pct']:.0f}%  ({f['n_upgrades']} upgrades, {f['n_downgrades']} downgrades)")
    print(f"    Grade changed:     {f['grade_change_pct']:.0f}%")
    print(f"    Either changed:    {f['either_change_pct']:.0f}%")
    print()
    print("  TOP CDU TRANSITIONS (what appraisers actually did):")
    for t, n in f["cdu_transitions"][:6]:
        print(f"    {t}: {n}x")
    if f["grade_transitions"]:
        print()
        print("  TOP GRADE TRANSITIONS:")
        for t, n in f["grade_transitions"][:4]:
            print(f"    {t}: {n}x")
    print()
    print("  UPGRADE SIGNALS (ranked by predictive power):")
    shown = 0
    for s in f["upgrade_stats"]:
        if s["multiplier"] >= 1.5 and s["count"] >= 2:
            print(f"    \"{s['phrase']}\": {s['count']}/{f['n_upgrades']} upgrades  "
                  f"({s['rate']*100:.0f}% vs {s['baseline']*100:.0f}% baseline, {s['multiplier']:.1f}x)")
            shown += 1
        if shown >= 8:
            break
    if not shown:
        print("    (no phrases met threshold this month)")
    print()
    print("  DOWNGRADE SIGNALS (ranked by predictive power):")
    shown = 0
    for s in f["downgrade_stats"]:
        if s["multiplier"] >= 1.5 and s["count"] >= 2:
            print(f"    \"{s['phrase']}\": {s['count']}/{f['n_downgrades']} downgrades  "
                  f"({s['rate']*100:.0f}% vs {s['baseline']*100:.0f}% baseline, {s['multiplier']:.1f}x)")
            shown += 1
        if shown >= 8:
            break
    if not shown:
        print("    (no phrases met threshold this month)")
    print()
    print(f"  HARD NEGATIVE LANGUAGE:")
    print(f"    Change rate: {f['hard_neg_change_pct']:.0f}%  (n={f['hard_neg_n']})")
    print()
    print(f"  LOCATION/SCHOOL FLUFF (noise):")
    print(f"    In changed parcels:   {f['noise_chg_pct']:.0f}%")
    print(f"    In unchanged parcels: {f['noise_unch_pct']:.0f}%")
    if f["yrblt_median_upgraded"]:
        print()
        print(f"  MEDIAN YEAR BUILT of upgraded parcels: {f['yrblt_median_upgraded']}")
    print("=" * 65)


# ── Step 7: Build and inject updated prompt ────────────────────────────────────

def build_prompt(f: dict) -> str:
    transitions_str = ", ".join(f"{t} ({n}x)" for t, n in f["cdu_transitions"][:6])

    upgrade_lines = []
    for s in f["upgrade_stats"]:
        if s["multiplier"] >= 1.5 and s["count"] >= 2:
            upgrade_lines.append(
                f'- "{s["phrase"]}" → {s["count"]}/{f["n_upgrades"]} upgrades '
                f'({s["rate"]*100:.0f}%, {s["multiplier"]:.1f}x baseline)'
            )
    downgrade_lines = []
    for s in f["downgrade_stats"]:
        if s["multiplier"] >= 1.5 and s["count"] >= 2:
            downgrade_lines.append(
                f'- "{s["phrase"]}" → {s["count"]}/{f["n_downgrades"]} downgrades '
                f'({s["rate"]*100:.0f}%, {s["multiplier"]:.1f}x baseline)'
            )

    upgrade_block   = "\n".join(upgrade_lines)   or "- (insufficient data this month)"
    downgrade_block = "\n".join(downgrade_lines) or "- (insufficient data this month)"
    trained_month   = date.today().strftime("%B %Y")

    return f"""const textPrompt = `You are a certified property appraiser assistant for the Stark County Auditor's Office in Ohio. Suggest a Grade and Condition based on all available evidence.

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
GRADE SCALE (2024 Stark County)
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
X+/X/X-  Excellent   — architect-designed, exceptional custom materials
A+/A/A-  Very Good   — high-quality, custom ornamentation, premium finishes
B+/B/B-  Good        — above-average, upgraded materials
C+/C/C-  Average     — mass-produced, meets code, stock materials
D        Fair        — below-average materials, plain interior
E        Low         — basic minimum code

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
CONDITION SCALE (1–5)
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
1 Excellent  — near new, 10–12 pluses
2 Good       — well maintained, 4–9 pluses
3 Average    — normal wear, 0–3 pluses/minuses
4 Fair       — deferred maintenance, 4–9 minuses
5 Poor       — major repairs needed, 10–12 minuses

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
STARK COUNTY EMPIRICAL DATA — {f['total']} PARCELS — Trained {trained_month}
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
These are ACTUAL appraiser decisions from this office. Weight heavily.

CHANGE RATES:
- Appraisers changed CDU on {f['cdu_change_pct']:.0f}% of parcels, Grade on {f['grade_change_pct']:.0f}%
- {f['either_change_pct']:.0f}% had at least one change; {100-f['either_change_pct']:.0f}% confirmed unchanged
- Most common CDU changes: {transitions_str}
- When uncertain between two CDU values lean toward GD or AV — these are the most common outcomes

HARD NEGATIVE LANGUAGE = STRONGEST SIGNAL:
- "Estate sale / cash only / sold as-is / vacant / abandoned" → {f['hard_neg_change_pct']:.0f}% CDU change rate (n={f['hard_neg_n']})
- These phrases = almost certainly FR or PR. Do not soften this.

CDU UPGRADE SIGNALS (ranked by Stark County predictive power):
{upgrade_block}

CDU DOWNGRADE SIGNALS (ranked by Stark County predictive power):
{downgrade_block}

YEAR BUILT: Homes with median YRBLT {f['yrblt_median_upgraded']} CAN be upgraded if truly renovated — age alone does not determine CDU.

LOCATION/SCHOOL FLUFF — IGNORE COMPLETELY:
"Great location / award-winning schools / minutes from shopping / desirable / sought-after"
appeared at {f['noise_chg_pct']:.0f}% in changed vs {f['noise_unch_pct']:.0f}% in unchanged parcels — zero predictive value.

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
EVIDENCE WEIGHTING HIERARCHY
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
1. PHOTO — highest weight. Roof, siding, windows, overall upkeep. If photo shows deferred maintenance, trust it over ANY positive remarks.
2. HARD NEGATIVE LANGUAGE — nearly always FR/PR regardless of other language
3. SPECIFIC NAMED SYSTEM UPDATES — remodeled/renovated/new furnace/new windows/updated kitchen+bath in combination
4. SOFT NEGATIVE LANGUAGE — opportunity/investor/as-is = strong lean toward FR
5. POSITIVE GENERIC LANGUAGE — lowest weight. "Beautifully updated/move-in ready" appeared on AV and FR properties. Do NOT use alone.
6. COSMETIC ONLY — fresh paint/new carpet/stainless/hardwood alone = no CDU change warranted

PHOTO vs REMARKS CONFLICT: If photo shows worn exterior but remarks claim "move-in ready" — TRUST THE PHOTO. Flag the conflict in rationale.

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
PROPERTY DATA
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
Address: \${{parcel.address}}, \${{parcel.city}} OH \${{parcel.zip}}
Year Built: \${{parcel.yrblt || "unknown"}}
Above-Grade Sq Ft (SFLA): \${{parcel.sfla != null ? parcel.sfla : "unknown"}}
\${{bathStr ? \`Bathrooms: \${{bathStr}}\` : ""}}
Current CAMA Grade: \${{parcel.grade || "unknown"}}
Current CAMA CDU: \${{parcel.cdu || "unknown"}}
MLS Public Remarks: \${{parcel.remarks || "(none provided)"}}
\${{hasPhoto ? "Front exterior photo attached — weight this heavily." : "No photo — reduce confidence, rely on remarks and year built only."}}

Respond ONLY with valid JSON — no markdown, no extra text:
{{
  "suggestedGrade": "e.g. B+",
  "suggestedCondition": 3,
  "conditionLabel": "e.g. Average",
  "gradeConfidence": "Low|Moderate|High",
  "conditionConfidence": "Low|Moderate|High",
  "rationale": "2–3 sentences. Cite specific evidence used, flag realtor bias detected, note any photo vs. remarks conflict."
}}`"""


def update_template(new_prompt: str) -> bool:
    html = TEMPLATE.read_text(encoding="utf-8")
    start = html.find("const textPrompt")
    if start == -1:
        print("ERROR: Could not find 'const textPrompt' in template")
        return False
    # Replace ONLY the template literal, up to its matching CLOSING backtick.
    # DO NOT search for the next "`;" — the prompt ends with a bare backtick
    # (followed by `try {`), so the next "`;" is ~79KB downstream and using it
    # deletes the entire Anthropic API-call block (blank-screen bug, 2026-07-06).
    open_bt = html.find("`", start)
    if open_bt == -1:
        print("ERROR: Could not find opening backtick of prompt")
        return False
    k = open_bt + 1
    while k < len(html) and not (html[k] == "`" and html[k - 1] != "\\"):
        k += 1
    if k >= len(html):
        print("ERROR: Could not find closing backtick of prompt")
        return False
    end = k + 1  # include the closing backtick; preserve everything after it
    TEMPLATE.write_text(html[:start] + new_prompt + html[end:], encoding="utf-8")
    return True


# ── Step 8: Rebuild portal ─────────────────────────────────────────────────────

def rebuild_portal():
    """Rebuild review_portal.html using this week's data."""
    import importlib.util, datetime as _dt

    # Find the week folder with the most recently dated value_mismatches file
    mlscama = ROOT / "MLSvsCAMA"
    best_folder = None
    best_date = None
    for f in mlscama.iterdir():
        if not f.is_dir():
            continue
        xlsx_files = list(f.glob("value_mismatches_*.xlsx"))
        if not xlsx_files:
            continue
        date_str_try = xlsx_files[0].stem.replace("value_mismatches_", "")
        try:
            d = _dt.date.fromisoformat(date_str_try)
            if best_date is None or d > best_date:
                best_date = d
                best_folder = f
                best_xlsx = xlsx_files[0]
        except ValueError:
            continue

    if best_folder is None:
        print("  WARNING: No value_mismatches file found in any week folder — skipping portal rebuild")
        return

    week_folder = best_folder
    print(f"  Using data from: {week_folder.name}")

    date_str = best_date.isoformat()
    xlsx_files = [best_xlsx]
    run_date = _dt.date.fromisoformat(date_str)

    spec = importlib.util.spec_from_file_location("build_portal", BUILD)
    bp   = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(bp)

    mismatches_rows = bp.xlsx_to_json(week_folder / f"value_mismatches_{date_str}.xlsx")
    perfects_rows   = bp.xlsx_to_json(week_folder / f"perfect_matches_{date_str}.xlsx")
    portal_photos   = week_folder / "Photos_New_Portal"
    photo_map       = bp.load_photos(portal_photos) if portal_photos.is_dir() else {}

    monday     = run_date - _dt.timedelta(days=run_date.weekday())
    week_label = f"Week of {monday.strftime('%B %d, %Y')}"
    week_key   = bp.get_iso_week_key(run_date)

    # One build marker shared by the page and version.json so staff on a cached
    # build get the "Reload now" prompt (mirrors build_portal.py).
    generated_at = _dt.datetime.now().strftime('%B %d, %Y at %I:%M %p')

    html = bp.build_html(
        mismatches_rows, perfects_rows, photo_map, week_label, generated_at,
        api_key="",
        github_pages_base=bp.GITHUB_PAGES_BASE,
        shared_api_url=bp.SHARED_API_URL,
    )
    output = ROOT / "review_portal.html"
    output.write_text(html, encoding="utf-8")
    size_mb = output.stat().st_size / (1024 * 1024)
    print(f"  Built: review_portal.html  ({size_mb:.1f} MB)")

    version_path = ROOT / "version.json"
    version_path.write_text(json.dumps({
        "generatedAt": generated_at,
        "weekLabel":   week_label,
        "weekKey":     week_key,
    }), encoding="utf-8")
    print(f"  Wrote: version.json")

    archive_path = bp.save_archive(mismatches_rows, perfects_rows, week_key, week_label, ROOT)
    print(f"  Archive: {archive_path.name}")


# ── Step 9: Push to GitHub ─────────────────────────────────────────────────────

def push_github(trained_month: str):
    try:
        for fname in ["review_portal.html", "version.json", "Grade&Con_Compara.xlsx"]:
            try:
                subprocess.run(["git", "add", fname], cwd=str(ROOT), check=True)
            except subprocess.CalledProcessError:
                pass  # file may be gitignored or unchanged
        subprocess.run(
            ["git", "commit", "-m", f"Monthly AI retrain — {trained_month}"],
            cwd=str(ROOT), check=True
        )
        subprocess.run(["git", "push"], cwd=str(ROOT), check=True)
        print(f"  Pushed to GitHub")
    except subprocess.CalledProcessError as e:
        print(f"  WARNING: GitHub push failed — push manually: {e}")


# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    trained_month = date.today().strftime("%B %Y")

    # Window is a flag, not a constant: a normal monthly run wants 4 weeks, but a
    # catch-up run after skipped months needs a wider one or those weeks' decisions
    # are silently discarded (load_recent_archives only ever takes the newest N).
    n_weeks = 4
    skip_recent = 0
    for i, a in enumerate(sys.argv):
        if a == "--weeks" and i + 1 < len(sys.argv):
            n_weeks = int(sys.argv[i + 1])
        elif a == "--skip-recent" and i + 1 < len(sys.argv):
            skip_recent = int(sys.argv[i + 1])

    print("\n" + "=" * 65)
    print("  STARK COUNTY — MONTHLY AI GRADE & CONDITION RETRAINING")
    print("=" * 65)

    print(f"\n[1/6] Loading recent archive data ({n_weeks} weeks)...")
    portal, weeks = load_recent_archives(n_weeks=n_weeks, skip_recent=skip_recent)

    print("\n[2/6] Querying Oracle for current Grade/CDU...")
    oracle = fetch_grade_cdu(list(portal.keys()))

    print("\n[3/6] Saving Grade&Con_Compara.xlsx...")
    save_compara(oracle)

    print("\n[4/6] Matching and analyzing appraiser decisions...")
    matched = build_matched(portal, oracle)
    print(f"  Matched: {len(matched)} parcels")

    if len(matched) < 30:
        print(f"\n  WARNING: Only {len(matched)} matched parcels — results may not be reliable.")
        ans = input("  Continue anyway? (y/n): ").strip().lower()
        if ans != "y":
            sys.exit(0)

    findings = analyze(matched)
    print_report(findings, weeks)

    print("\n[5/6] Updating AI prompt in template...")
    new_prompt = build_prompt(findings)
    if update_template(new_prompt):
        print(f"  Updated: review_portal_template.html")
    else:
        print("  ERROR: Failed to update template")
        sys.exit(1)

    prompt_only = "--prompt-only" in sys.argv
    if prompt_only:
        print("\n[6/6] --prompt-only: skipping portal rebuild and push.")
        print("  New prompt is saved in review_portal_template.html and will ship")
        print("  automatically with next week's build_portal.py run.")
    else:
        print("\n[6/6] Rebuilding portal and pushing to GitHub...")
        rebuild_portal()
        push_github(trained_month)

    print("\n" + "=" * 65)
    print("  RETRAINING COMPLETE")
    if prompt_only:
        print("  Prompt updated in template only — goes live with next week's build.")
    else:
        print(f"  Staff will see updated AI suggestions immediately.")
    print("=" * 65)

    try:
        input("\nPress Enter to close...")
    except EOFError:
        pass


if __name__ == "__main__":
    main()
